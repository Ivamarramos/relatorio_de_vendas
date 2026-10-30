import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ARQUIVO_ENTRADA = "vendas.csv"
ARQUIVO_SAIDA = "relatorio_vendas.xlsx"
COLUNAS_MOEDA = ["preco_unitario", "valor_total", "faturamento"]

def formatar_reais(valor):
    texto = f"{valor:,.2f}"
    texto = texto.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {texto}"

def carregar_dados(caminho):
    df = pd.read_csv(caminho, parse_dates=["data"])
    df["valor_total"] = df["quantidade"] * df["preco_unitario"]
    df["mes"] = df["data"].dt.strftime("%Y-%m")
    return df.sort_values("data")

def somar_por(df, coluna):
    return (
        df.groupby(coluna)
        .agg(quantidade=("quantidade", "sum"), faturamento=("valor_total", "sum"))
        .reset_index()
    )
    
def montar_resumo(df, por_produto, por_mes):
    melhor_mes = por_mes.loc[por_mes["faturamento"].idxmax(), "mes"]
    indicadores = [
        ("Total de vendas (registro)", len(df)),
        ("Faturamento total", formatar_reais(df["valor_total"].sum())),
        ("Ticket médio", formatar_reais(df["valor_total"].mean())),
        ("Produto com maior faturamento", por_produto.iloc[0]["produto"]),
        ("Melhor mês", melhor_mes),
    ]       
    
    return pd.DataFrame(indicadores, columns=["Indicador", "Valor"])

def formatar_planilha(planilha):
    for celula in planilha[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.alignment = Alignment(horizontal="center")
        
        cabecalhos = [celula.value for celula in planilha[1]]
        for posicao, nome in enumerate(cabecalhos, start=1):
            if nome in COLUNAS_MOEDA:
                for linha in range(2, planilha.max_row + 1):
                    planilha.cell(row=linha, column=posicao).number_format = "R$ #,##0.00"
                    
                for coluna in planilha.columns:
                    maior = max(len(str(celula.value)) if celula.value is not None else 0 for celula in coluna)
                    letra = get_column_letter(coluna[0].column)
                    planilha.column_dimensions[letra].width = maior + 4 
                    
def gerar_excel(caminho, abas):
    with pd.ExcelWriter(caminho, engine="openpyxl") as escritor:
        for nome, tabela in abas.items():
            tabela.to_excel(escritor, sheet_name=nome, index=False)
            formatar_planilha(escritor.sheets[nome])

def main():
    try:
        df = carregar_dados(ARQUIVO_ENTRADA)
    except FileNotFoundError:
        print("Arquivo vendas.csv não encontrado. Rode primeiro: python gerar_dados.py")
        return
    
    por_produto = somar_por(df, "produto").sort_values("faturamento", ascending=False)
    por_mes = somar_por(df, "mes")
    por_vendedor = somar_por(df, "vendedor").sort_values("faturamento", ascending=False)
    
    dados = df.copy()
    dados["data"] = dados["data"].dt.date
    
    abas = {
        "Resumo": montar_resumo(df, por_produto, por_mes),
        "Por Produto": por_produto,
        "Por Mês": por_mes,
        "Por Vendedor": por_vendedor,
        "Dados": dados,
    }
    
    try:
        gerar_excel(ARQUIVO_SAIDA, abas)
    except PermissionError:
        print("Feche o arquivo relatorio_vendas.xksx no Excel e tente de novo.")
        return
    
    print(f"Relatório gerado: {ARQUIVO_SAIDA}")
    
if __name__ == "__main__":
    main()

                    